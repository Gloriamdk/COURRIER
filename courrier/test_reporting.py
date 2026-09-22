from datetime import datetime, timedelta
from io import BytesIO

from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone

from .models import Affectation, Courrier, Decision, Historique, ReponseCourrier, User
from .reporting import ReportFilters, build_report


class ReportingTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.now = timezone.make_aware(datetime(2026, 6, 15, 12))
        cls.users = {role: User.objects.create_user(username='report_' + role, role=role)
                     for role in User.Role.values}
        cls.sg = cls.users[User.Role.SG]

    def mail(self, reference, status=Courrier.Statut.AFFECTE, priority=Courrier.Priorite.NORMAL, arrived=None):
        return Courrier.objects.create(reference=reference, designation='Objet ' + reference,
            expediteur_nom='Expéditeur', cree_par=self.sg, statut=status, priorite=priority,
            date_arrivee=arrived or self.now - timedelta(days=15), delai_traitement_jours=10)

    def assignment(self, mail, direction='DAAF', start=None, deadline=None, finished=None, recipient=None):
        decision, _ = Decision.objects.get_or_create(courrier=mail, defaults={'signe_par': self.users[User.Role.MINISTRE]})
        assignment = Affectation.objects.create(courrier=mail, decision=decision,
            affecte_par=self.sg, service_concerne=direction, destinataire=recipient,
            date_limite_traitement=deadline, date_traitement=finished,
            statut_traitement=Affectation.StatutTraitement.TRAITE if finished else Affectation.StatutTraitement.RECU)
        Affectation.objects.filter(pk=assignment.pk).update(date_affectation=start or self.now - timedelta(days=10))
        return assignment

    def report(self, **params):
        form = ReportFilters({'annee': '2026', **params})
        self.assertTrue(form.is_valid(), form.errors)
        return build_report(form.cleaned_data, now=self.now)

    def fixture(self):
        completed = self.mail('DONE', Courrier.Statut.TERMINE, Courrier.Priorite.URGENT)
        self.assignment(completed, 'DAF', deadline=self.now - timedelta(days=1), finished=self.now - timedelta(days=2))
        # Duplicate assignments and responses must not multiply the mail count.
        self.assignment(completed, 'DAAF', deadline=self.now - timedelta(days=1), finished=self.now - timedelta(days=2))
        for version in (1, 2):
            ReponseCourrier.objects.create(courrier=completed, auteur=self.sg, version=version)
            Historique.objects.create(courrier=completed, utilisateur=self.sg, action='NOTE', description='Note')
        late = self.mail('LATE', priority=Courrier.Priorite.TRES_URGENT)
        self.assignment(late, deadline=self.now - timedelta(days=1))
        self.assignment(late, deadline=self.now - timedelta(hours=1))
        pending = self.mail('PENDING', Courrier.Statut.VALIDE_DIRECTEUR)
        self.assignment(pending, 'DPDT', deadline=self.now + timedelta(days=2))
        no_deadline = self.mail('NO-DEADLINE', Courrier.Statut.ARRIVE)
        return completed, late, pending, no_deadline

    def test_three_roles_can_access_page_and_exports(self):
        for role in (User.Role.SG, User.Role.DC, User.Role.MINISTRE):
            self.client.force_login(self.users[role])
            for name in ('rapports_statistiques', 'rapports_pdf', 'rapports_excel'):
                with self.subTest(role=role, endpoint=name):
                    response = self.client.get(reverse(name), {'annee': 2026})
                    self.assertEqual(response.status_code, 200)
                    if name == 'rapports_statistiques':
                        self.assertContains(response, 'id="nav-rapports"')
                        self.assertContains(response, 'report-chart-data')
                    elif name == 'rapports_pdf':
                        self.assertTrue(response.content.startswith(b'%PDF-'))
                    else:
                        self.assertTrue(response.content.startswith(b'PK'))

    def test_every_other_role_denied_including_superuser_agent(self):
        for role, user in self.users.items():
            if role in (User.Role.SG, User.Role.DC, User.Role.MINISTRE):
                continue
            self.client.force_login(user)
            for name in ('rapports_statistiques', 'rapports_pdf', 'rapports_excel'):
                with self.subTest(role=role, endpoint=name):
                    self.assertEqual(self.client.get(reverse(name)).status_code, 403)
        agent = self.users[User.Role.AGENT]
        agent.is_superuser = True
        agent.save(update_fields=['is_superuser'])
        self.client.force_login(agent)
        self.assertEqual(self.client.get(reverse('rapports_statistiques')).status_code, 403)

    def test_anonymous_redirected_and_inactive_denied(self):
        self.assertEqual(self.client.get(reverse('rapports_statistiques')).status_code, 302)
        self.sg.is_active = False
        self.sg.save(update_fields=['is_active'])
        self.client.force_login(self.sg)
        self.assertEqual(self.client.get(reverse('rapports_statistiques')).status_code, 302)

    def test_menu_hidden_for_other_roles(self):
        from django.template.loader import render_to_string
        for role, user in self.users.items():
            html = render_to_string('base.html', {'user': user})
            self.assertEqual('id="nav-rapports"' in html, role in ('SG', 'DC', 'MINISTRE'))

    def test_counts_rate_urgencies_and_no_duplicates(self):
        self.fixture()
        report = self.report()
        self.assertEqual(report['kpi'], dict(total=4, treated=1, ongoing=3, late=1, urgent=2, rate=25.0, closed=1))
        self.assertEqual(report['urgent']['rate'], 50.0)
        self.assertEqual(report['urgent']['late'], 1)
        self.assertEqual(sum(r['count'] for r in report['distributions']['statut']), 4)
        self.assertEqual(sum(r['count'] for r in report['distributions']['priorite']), 4)

    def test_direction_alias_counts_and_filter(self):
        self.fixture()
        rows = self.report()['directions']
        self.assertEqual([(r['direction'], r['total'], r['treated'], r['late']) for r in rows], [('DAAF', 2, 1, 1), ('DPDT', 1, 0, 0)])
        report = self.report(direction='DAAF')
        self.assertEqual(report['kpi']['total'], 2)
        self.assertEqual(report['kpi']['rate'], 50.0)
        self.assertEqual(len(report['directions']), 1)

    def test_multiple_directions_and_scoped_deadlines(self):
        mail = self.mail('SHARED')
        self.assignment(mail, 'DAAF', deadline=self.now - timedelta(days=1))
        self.assignment(mail, 'DPDT', deadline=self.now + timedelta(days=1))
        report = self.report()
        self.assertEqual(report['kpi']['total'], 1)
        self.assertEqual(sum(r['total'] for r in report['directions']), 2)
        self.assertEqual(self.report(direction='DPDT')['kpi']['late'], 0)
        self.assertEqual(self.report(direction='DAAF')['kpi']['late'], 1)

    def test_recipient_direction_fallback_and_unknown(self):
        agent = self.users[User.Role.AGENT]
        agent.service_direction = ' DAF '
        agent.save(update_fields=['service_direction'])
        self.assignment(self.mail('FALLBACK'), '', recipient=agent)
        self.assignment(self.mail('UNKNOWN'), None)
        self.assertEqual({r['direction'] for r in self.report()['directions']}, {'DAAF', 'Sans direction renseignée'})

    def test_only_termine_is_completed(self):
        for status in Courrier.Statut.values:
            self.mail(status, status)
        self.assertEqual(self.report()['kpi']['treated'], 1)
        self.assertEqual(self.report()['kpi']['ongoing'], len(Courrier.Statut.values)-1)

    def test_delay_mean_median_and_deadline_compliance(self):
        for i, days in enumerate([2, 4, 9]):
            mail = self.mail('DURATION' + str(i), Courrier.Statut.TERMINE)
            self.assignment(mail, start=self.now - timedelta(days=10),
                deadline=self.now - timedelta(days=5), finished=self.now - timedelta(days=10-days))
        t = self.report()['timing']
        self.assertEqual(t['average'], 5)
        self.assertEqual(t['median'], 4)
        self.assertEqual(t['on_time'], 2)
        self.assertEqual(t['after_deadline'], 1)
        self.assertEqual(t['rate'], 66.7)
        self.assertEqual(self.report()['kpi']['late'], 0)

    def test_history_closure_takes_precedence_and_latest_transition(self):
        mail = self.mail('HISTORY', Courrier.Statut.TERMINE)
        self.assignment(mail, finished=self.now - timedelta(days=8))
        for days in (4, 2):
            h = Historique.objects.create(courrier=mail, action='CLOTURE', description='Clôture', nouveau_statut='TERMINE', ancien_statut='AFFECTE')
            Historique.objects.filter(pk=h.pk).update(date_action=self.now - timedelta(days=days))
        self.assertEqual(self.report()['timing']['average'], 8)

    def test_missing_and_negative_dates_not_fabricated(self):
        self.mail('NO-DATE', Courrier.Statut.TERMINE)
        self.assignment(self.mail('NEGATIVE', Courrier.Statut.TERMINE),
            start=self.now, finished=self.now - timedelta(days=1))
        report = self.report()
        self.assertIsNone(report['timing']['average'])
        self.assertIsNone(report['timing']['median'])
        self.assertEqual(report['timing']['missing'], 2)
        self.assertEqual(report['unknown_finish'], 1)

    def test_late_matches_existing_alert_rules(self):
        mail = self.mail('NO-DELAY')
        self.assignment(mail, deadline=self.now - timedelta(days=1))
        mail.delai_traitement_jours = None
        mail.save(update_fields=['delai_traitement_jours'])
        self.assignment(self.mail('AT-LIMIT'), deadline=self.now)
        self.assignment(self.mail('FINISHED-ASSIGNMENT'), deadline=self.now - timedelta(days=1), finished=self.now)
        self.assertEqual(self.report()['kpi']['late'], 0)

    def test_date_range_includes_whole_end_day(self):
        self.mail('START', arrived=timezone.make_aware(datetime(2026, 5, 1)))
        self.mail('END', arrived=timezone.make_aware(datetime(2026, 5, 31, 23, 59, 59)))
        self.mail('OUTSIDE', arrived=timezone.make_aware(datetime(2026, 6, 1)))
        self.assertEqual(self.report(debut='2026-05-01', fin='2026-05-31')['kpi']['total'], 2)
        self.assertEqual(self.report(mois='5')['kpi']['total'], 2)

    def test_explicit_dates_do_not_implicitly_restrict_current_year(self):
        form = ReportFilters({'debut': '2024-01-01', 'fin': '2025-12-31'})
        self.assertTrue(form.is_valid())
        self.assertIsNone(form.cleaned_data['annee'])

    def test_default_is_current_year(self):
        form = ReportFilters({})
        self.assertTrue(form.is_valid())
        self.assertEqual(form.cleaned_data['start'].year, timezone.localdate().year)

    def test_invalid_filters_return_400_also_for_exports(self):
        self.client.force_login(self.sg)
        for params in [{'debut': 'invalid'}, {'debut': '2026-02-01', 'fin': '2026-01-01'},
                       {'mois': '13'}, {'annee': '99999'}, {'direction': 'does-not-exist'},
                       {'statut': 'FICTIF'}, {'annee': '2026', 'debut': '2025-01-01', 'fin': '2025-02-01'}]:
            for endpoint in ('rapports_statistiques', 'rapports_pdf', 'rapports_excel'):
                self.assertEqual(self.client.get(reverse(endpoint), params).status_code, 400)

    def test_status_priority_and_urgent_filters(self):
        self.fixture()
        self.assertEqual(self.report(statut='TERMINE')['kpi']['total'], 1)
        self.assertEqual(self.report(priorite='TRES_URGENT')['kpi']['total'], 1)
        self.assertEqual(self.report(urgent='oui')['kpi']['total'], 2)
        self.assertEqual(self.report(urgent='non')['kpi']['total'], 2)

    def test_zero_mail_and_single_direction(self):
        report = self.report()
        self.assertEqual(report['kpi']['total'], 0)
        self.assertEqual(report['kpi']['rate'], 0)
        self.assertEqual(report['directions'], [])
        self.assertIsNone(report['comparison']['rows'][0]['change'])
        self.assignment(self.mail('SINGLE'))
        self.assertEqual(len(self.report()['directions']), 1)

    def test_monthly_uses_real_closure_date(self):
        mail = self.mail('MONTH', Courrier.Statut.TERMINE, arrived=timezone.make_aware(datetime(2026, 1, 1)))
        self.assignment(mail, finished=timezone.make_aware(datetime(2026, 2, 1)))
        rows = {r['month']: r for r in self.report()['monthly']}
        self.assertEqual(rows['2026-01']['received'], 1)
        self.assertEqual(rows['2026-01']['treated'], 0)
        self.assertEqual(rows['2026-02']['treated'], 1)

    def test_comparison_and_zero_denominator(self):
        self.mail('PREVIOUS', arrived=timezone.make_aware(datetime(2025, 5, 1)))
        self.mail('CURRENT-1')
        self.mail('CURRENT-2')
        rows = self.report()['comparison']['rows']
        self.assertEqual(rows[0]['previous'], 1)
        self.assertEqual(rows[0]['change'], 100.0)
        self.assertIsNone(rows[1]['change'])

    def test_exports_preserve_filters_and_escape_spreadsheet_formulas(self):
        from openpyxl import load_workbook
        self.assignment(self.mail('FORMULA'), '=1+1')
        self.fixture()
        self.client.force_login(self.sg)
        response = self.client.get(reverse('rapports_excel'), {'annee': 2026, 'direction': '=1+1'})
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook['Directions']['A2'].value, '=1+1')
        self.assertEqual(workbook['Directions']['A2'].data_type, 's')
        self.assertEqual(workbook['Indicateurs']['B2'].value, 1)
        self.assertEqual(workbook['Directions'].max_row, 2)
        response = self.client.get(reverse('rapports_pdf'), {'annee': 2026, 'direction': 'DAAF'})
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertGreater(len(response.content), 5000)

    def test_queries_are_read_only_and_do_not_grow_per_direction(self):
        self.fixture()
        with CaptureQueriesContext(connection) as first:
            self.report()
        for i in range(12):
            self.assignment(self.mail('EXTRA' + str(i)), 'Direction ' + str(i))
        with CaptureQueriesContext(connection) as second:
            self.report()
        self.assertEqual(len(first), len(second))
        self.assertLessEqual(len(second), 16)
        self.assertTrue(all(q['sql'].lstrip().upper().startswith('SELECT') for q in second))

    def test_json_script_escapes_untrusted_direction(self):
        self.assignment(self.mail('XSS'), '</script><script>alert(1)</script>')
        self.client.force_login(self.sg)
        response = self.client.get(reverse('rapports_statistiques'), {'annee': 2026})
        self.assertNotContains(response, '</script><script>alert(1)</script>')
        self.assertContains(response, r'\u003C/script\u003E')
