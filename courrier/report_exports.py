"""Server-generated exports share the exact dashboard calculations and filters."""
from io import BytesIO
from html import escape

from django.conf import settings
from django.utils import timezone


def period_label(report):
    f = report['filters']
    return f"Réception : {f['start'] or 'sans borne de début'} au {f['end'] or 'sans borne de fin'}"


def report_sections(report):
    k, t, u = report['kpi'], report['timing'], report['urgent']
    return [
        ('Indicateurs', [['Indicateur', 'Valeur'], ['Courriers reçus', k['total']],
            ['Traités / clôturés', k['treated']], ['En cours (tous les non terminés)', k['ongoing']],
            ['Actuellement en retard', k['late']], ['Urgents', k['urgent']],
            ['Confidentiels', 'Non disponible'], ['Taux de traitement', f"{k['rate']} %"]]),
        ('Directions', [['Direction', 'Affectés', 'Traités', 'En cours', 'En retard', 'Taux (%)']] +
            [[r[x] for x in ('direction', 'total', 'treated', 'ongoing', 'late', 'rate')] for r in report['directions']]),
        ('Délais', [['Indicateur', 'Valeur'],
            ['Moyenne (jours)', round(t['average'], 2) if t['average'] is not None else 'Non disponible'],
            ['Médiane (jours)', round(t['median'], 2) if t['median'] is not None else 'Non disponible'],
            ['Courriers mesurables', t['measured']], ['Dates absentes ou incohérentes', t['missing']],
            ['Traités dans les délais', t['on_time']], ['Traités après échéance', t['after_deadline']],
            ['Actuellement en retard', k['late']], ['Respect des délais', f"{t['rate']} %"]]),
        ('Urgences', [['Indicateur', 'Valeur'], ['Total', u['total']], ['Traités', u['treated']],
            ['En cours', u['ongoing']], ['En retard', u['late']], ['Taux de traitement', f"{u['rate']} %"]]),
        ('Évolution mensuelle', [['Mois', 'Reçus', 'Traités']] +
            [[r['month'], r['received'], r['treated']] for r in report['monthly']]),
        ('Statuts', [['Statut', 'Nombre', 'Part (%)']] +
            [[r['label'], r['count'], r['rate']] for r in report['distributions']['statut']]),
        ('Priorités', [['Priorité', 'Nombre', 'Part (%)', 'En retard']] +
            [[r['label'], r['count'], r['rate'], r['late']] for r in report['distributions']['priorite']]),
    ]


METHODOLOGY = (
    "Périmètre : courriers entrants sélectionnés par date de réception. Traité et clôturé : statut TERMINE. "
    "Tous les autres statuts restent en cours. Les courriers sortants sont des réponses et ne sont pas additionnés aux entrants. "
    "Une direction est issue du service de l’affectation, à défaut du service actuel du destinataire ; DAF est regroupé avec DAAF. "
    "Un courrier peut figurer dans plusieurs directions ; leurs totaux ne s’additionnent pas au total global. "
    "Fin : dernière transition historique vers TERMINE, à défaut dernière finalisation d’affectation. "
    "Durée : première affectation du périmètre à cette fin. Les dates manquantes ou inversées sont exclues des durées. "
    "Échéance : première limite enregistrée des affectations du périmètre, si un délai propre au courrier est défini. "
    "Retard actuel : courrier non terminé avec au moins une affectation reçue/en cours dont la limite est dépassée. "
    "Respect des délais : traités à temps / traités avec date de fin et échéance connues. "
    "Évolution : dates de réception et de clôture des courriers du périmètre, y compris les clôtures ultérieures. "
    "Comparaison : mêmes dates de réception un an auparavant, états observés à la génération, et non états historiques. "
    "La confidentialité n’est pas enregistrée dans le modèle actuel."
)


def export_pdf(report):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image
    from reportlab.graphics.shapes import Drawing, String, Rect

    output = BytesIO()
    styles = getSampleStyleSheet()
    styles['BodyText'].fontSize = 9
    styles['BodyText'].leading = 13
    story = []
    logo = settings.BASE_DIR / 'static' / 'images' / 'logo_gec.png'
    if logo.exists():
        from PIL import Image as PILImage
        with PILImage.open(logo) as source:
            width, height = source.size
        story.append(Image(str(logo), width=95, height=95 * height / width))
    def paragraph(value, style='BodyText'):
        return Paragraph(escape(str(value)), styles[style])
    story.extend([paragraph('Ministère du Tourisme, de la Culture et des Arts', 'Heading2'),
        paragraph('Rapport statistique de gestion des courriers', 'Title'),
        paragraph(period_label(report)),
        paragraph('Généré le ' + timezone.localtime(report['generated_at']).strftime('%d/%m/%Y à %H:%M')),
        paragraph('Filtres : ' + ', '.join(f'{label} : {value}' for label, value in report['filter_labels']))])
    for title, rows in report_sections(report):
        story.extend([Spacer(1, 12), paragraph(title, 'Heading2')])
        widths = [170] + [(345 / (len(rows[0])-1))] * (len(rows[0])-1)
        table = Table([[paragraph(cell) for cell in row] for row in rows], colWidths=widths, repeatRows=1, hAlign='LEFT')
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#dceee8')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f3f6f8')]),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('BOTTOMPADDING', (0, 0), (-1, -1), 7),
            ('TOPPADDING', (0, 0), (-1, -1), 7)]))
        story.append(table)
    # Vector charts generated on the server also work when JavaScript is unavailable.
    for title, rows in [('Volume par direction (20 premières)', report['directions'][:20]),
                         ('Évolution mensuelle (24 derniers mois)', report['monthly'][-24:])]:
        if not rows:
            continue
        story.extend([Spacer(1, 12), paragraph(title, 'Heading2')])
        drawing = Drawing(510, len(rows)*22+30)
        maximum = max([r.get('total', r.get('received', 0)) for r in rows] + [r['treated'] for r in rows] + [1])
        for i, row in enumerate(rows):
            y = (len(rows)-i)*22
            label = row.get('direction', row.get('month', ''))
            total = row.get('total', row.get('received', 0))
            drawing.add(String(0, y, label[:27], fontSize=8))
            drawing.add(Rect(145, y, total*285/maximum, 7, fillColor=colors.HexColor('#3675a9'), strokeColor=None))
            drawing.add(Rect(145, y-8, row['treated']*285/maximum, 7, fillColor=colors.HexColor('#198268'), strokeColor=None))
            drawing.add(String(440, y, f"{total} / {row['treated']}", fontSize=8))
        drawing.add(String(145, 0, 'Bleu : reçus / affectés ; vert : traités', fontSize=8))
        story.append(drawing)
    if report['comparison']:
        c = report['comparison']
        story.append(paragraph(f"Comparaison au {c['start']} – {c['end']}", 'Heading2'))
        for r in c['rows']:
            change = f"{r['change']} %" if r['change'] is not None else 'non calculable (base nulle)'
            story.append(paragraph(f"{r['label']} : {r['current']} contre {r['previous']} ; évolution {change}."))
    story.append(paragraph('Synthèse', 'Heading2'))
    story.extend(paragraph(line) for line in report['synthesis'])
    story.extend([paragraph('Définitions et limites', 'Heading2'), paragraph(METHODOLOGY)])
    def footer(canvas, doc):
        canvas.setFont('Helvetica', 8)
        canvas.drawRightString(555, 22, f'MTCA · Page {doc.page}')
    SimpleDocTemplate(output, pagesize=A4, rightMargin=40, leftMargin=40,
                      topMargin=30, bottomMargin=38).build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def export_excel(report):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    workbook = Workbook()
    info = workbook.active
    info.title = 'Rapport'
    for row in [['Ministère du Tourisme, de la Culture et des Arts'],
                ['Rapport statistique de gestion des courriers'], [period_label(report)],
                ['Généré le', timezone.localtime(report['generated_at']).isoformat()],
                *report['filter_labels'], *[[s] for s in report['synthesis']], [METHODOLOGY]]:
        info.append(row)
    for title, rows in report_sections(report):
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'  # Never interpret user-provided direction names as formulas.
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill('solid', fgColor='DCEEE8')
        sheet.column_dimensions['A'].width = 48
        for key in 'BCDEF':
            sheet.column_dimensions[key].width = 20
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
